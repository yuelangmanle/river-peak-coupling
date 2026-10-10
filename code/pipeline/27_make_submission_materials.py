#!/usr/bin/env python3
"""Build clean editable Elsevier upload materials from the maintained text files."""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

BASE = Path(__file__).resolve().parents[2]
PAPER = BASE / "paper"


def normalize(document):
    for style in document.styles:
        try:
            style.font.name = "Times New Roman"
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.size = Pt(11)
        except AttributeError:
            pass
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.size = Pt(11)


def build_highlights():
    lines = [
        line.strip().lstrip("- ")
        for line in (PAPER / "joh_highlights.txt").read_text().splitlines()
        if line.strip() and not line.strip().startswith("#")
    ]
    doc = Document()
    for line in lines:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(line)
    normalize(doc)
    doc.save(PAPER / "Highlights.docx")


def build_dci():
    doc = Document()
    doc.add_paragraph("The authors declare no competing interests.")
    normalize(doc)
    doc.save(PAPER / "Declaration_of_Competing_Interests.docx")


def build_cover_letter():
    text = (PAPER / "Cover_Letter_JoH.txt").read_text().strip().splitlines()
    doc = Document()
    for line in text:
        if not line.strip():
            doc.add_paragraph()
        else:
            doc.add_paragraph(line)
    normalize(doc)
    doc.save(PAPER / "Cover_Letter_JoH.docx")


if __name__ == "__main__":
    build_highlights()
    build_dci()
    build_cover_letter()
    print("clean Highlights.docx, Declaration_of_Competing_Interests.docx and Cover_Letter_JoH.docx written")
